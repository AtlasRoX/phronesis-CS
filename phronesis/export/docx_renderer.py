"""DOCX renderer for the investigation report (owned by EXPORT, feature 9). python-docx is imported lazily."""

from __future__ import annotations

import io
from typing import Any, List

from .markdown_blocks import Block, Run, parse_blocks
from .pdf_renderer import BRAND_MARK, BRAND_WORDMARK
from .report_model import (
    ReportModel, citation_line, fmt_number, question_is_title, report_headline, report_kind, task_lines,
    verification_lines,
)

# The web app's typeface (bundled for the PDF under fonts/, SIL OFL; free from Google Fonts for Word). Word cannot
# embed it or fall back per font, so code uses Consolas, which every Office install has; the PDF uses Google Sans Code.
_FONT = "Google Sans"
_CODE_FONT = "Consolas"
_INK = (0x20, 0x21, 0x24)
_NAVY = (0x18, 0x1E, 0x48)
_MUTED = (0x5F, 0x63, 0x68)


def _add_runs(par: Any, runs: List[Run], *, force_bold: bool = False) -> None:
    from docx.shared import Pt

    for r in runs:
        if not r.text:
            continue
        run = par.add_run(r.text)
        run.bold = r.bold or force_bold or None
        run.italic = r.italic or None
        if r.strike:
            run.font.strike = True
        if r.code:
            run.font.name = _CODE_FONT
            run.font.size = Pt(9)
        if r.link:
            run.underline = True


def _list_style(block: Block) -> str:
    base = "List Number" if block.list_kind == "ordered" else "List Bullet"
    return base if block.depth <= 1 else f"{base} {min(block.depth, 3)}"


def _add_table(doc: Any, header: List[Any], rows: List[List[Any]], *, runs_mode: bool) -> None:
    from docx.shared import Pt

    spacer = doc.add_paragraph()  # tables take no space before them; keep them off the text above
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)
    ncols = max([len(header)] + [len(r) for r in rows] + [1])
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    if header:
        cells = table.add_row().cells
        for i, value in enumerate(header):
            par = cells[i].paragraphs[0]
            if runs_mode:
                _add_runs(par, value, force_bold=True)
            else:
                par.add_run(str(value)).bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row[:ncols]):
            par = cells[i].paragraphs[0]
            if runs_mode:
                _add_runs(par, value)
            else:
                par.add_run(str(value))
    _style_table(table, has_header=bool(header))


def _style_table(table: Any, *, has_header: bool) -> None:
    """Light grey rules, a tinted header row and compact 9 pt cells."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "DADCE0")
        borders.append(el)
    table._tbl.tblPr.append(borders)
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            if r == 0 and has_header:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:fill"), "F1F3F8")
                cell._tc.get_or_add_tcPr().append(shading)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.1
                for run in par.runs:
                    run.font.size = Pt(9)


def add_markdown(doc: Any, text: str, *, base_level: int = 2) -> None:
    from docx.shared import Pt

    blocks = parse_blocks(text)
    levels = [b.level for b in blocks if b.type == "heading"]
    top = min(levels) if levels else 1
    for b in blocks:
        if b.type == "heading":
            par = doc.add_heading("", level=min(base_level + (b.level - top), 4))
            _add_runs(par, b.runs)
        elif b.type == "paragraph":
            if b.list_kind:
                style = _list_style(b) if b.first_in_item else None
                par = doc.add_paragraph(style=style)
            elif b.quote:
                par = doc.add_paragraph(style="Quote")
            else:
                par = doc.add_paragraph()
            _add_runs(par, b.runs)
        elif b.type == "table":
            _add_table(doc, b.header, b.rows, runs_mode=True)
            doc.add_paragraph()
        elif b.type == "code":
            par = doc.add_paragraph()
            run = par.add_run(b.text)
            run.font.name = _CODE_FONT
            run.font.size = Pt(9)
        elif b.type == "hr":
            doc.add_paragraph("―" * 20)


def _set_font(style: Any, name: str) -> None:
    """Point every script slot at the font and drop the theme-font references that would override it."""
    from docx.oxml.ns import qn

    style.font.name = name
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(attr), None)


def _style_document(doc: Any) -> None:
    """Google Sans everywhere, readable line height and consistent spacing (mirrors the PDF report)."""
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    styles = doc.styles
    normal = styles["Normal"]
    _set_font(normal, _FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(*_INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)  # the template's accent-coloured rule under the title
    for name, size, before, after, color in (
        ("Title", 20, 0, 10, _NAVY),
        ("Heading 1", 14, 16, 4, _NAVY),
        ("Heading 2", 12, 12, 3, _NAVY),
        ("Heading 3", 11, 10, 2, _NAVY),
        ("Heading 4", 10.5, 8, 2, _NAVY),
    ):
        style = styles[name]
        _set_font(style, _FONT)
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.italic = False
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number", "List Number 2", "List Number 3"):
        styles[name].paragraph_format.space_after = Pt(3)
    quote = styles["Intense Quote"]
    _set_font(quote, _FONT)
    quote.font.color.rgb = RGBColor(0x92, 0x54, 0x00)
    caption = styles["Caption"] if "Caption" in [s.name for s in styles] else None
    if caption is not None:
        _set_font(caption, _FONT)
        caption.font.color.rgb = RGBColor(*_MUTED)


def _field(par: Any, code: str) -> Any:
    """A Word field (PAGE, NUMPAGES) that Word fills in when the document is opened or printed."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = par.add_run()
    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == code:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {code} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)
    return run


def _right_tab_only(par: Any, width: Any) -> None:
    """One right-aligned tab at the text edge; clears the Header/Footer styles' centre tab that would catch it first."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Inches

    stops = par.paragraph_format.tab_stops
    stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)


def _brand_header_footer(doc: Any, report: ReportModel) -> None:
    """Header: the Phronesis mark and wordmark, the report kind on the right. Footer: run id and page x of y."""
    import os

    from docx.shared import Pt, RGBColor

    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    head = section.header.paragraphs[0]
    _right_tab_only(head, width)
    if os.path.isfile(BRAND_MARK) and os.path.isfile(BRAND_WORDMARK):
        head.add_run().add_picture(BRAND_MARK, height=Pt(14))
        head.add_run(" ")
        head.add_run().add_picture(BRAND_WORDMARK, height=Pt(9))
    else:
        brand = head.add_run("Phronesis")
        brand.bold = True
        brand.font.color.rgb = RGBColor(*_NAVY)
    kind = head.add_run("\t" + report_kind(report))
    kind.font.size = Pt(8)
    kind.font.color.rgb = RGBColor(*_MUTED)

    foot = section.footer.paragraphs[0]
    _right_tab_only(foot, width)
    runs = [foot.add_run("Phronesis · " + (f"Run {report.run_id}" if report.run_id else "No stored run") + "\tPage ")]
    runs.append(_field(foot, "PAGE"))
    runs.append(foot.add_run(" of "))
    runs.append(_field(foot, "NUMPAGES"))
    for run in runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(*_MUTED)


def _title_block(doc: Any, report: ReportModel) -> None:
    """Report kind, the question as the title, and the headline facts as a borderless grid (mirrors the PDF)."""
    from docx.shared import Pt, RGBColor

    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(2)
    run = eyebrow.add_run(report_kind(report).upper())
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(*_MUTED)
    question = (report.question or "").strip()
    title = question if question_is_title(report) else report.title
    doc.add_paragraph(title[:1].upper() + title[1:], style="Title")

    facts = report_headline(report)
    cols = 3
    grid = doc.add_table(rows=0, cols=cols)
    for start in range(0, len(facts), cols):
        cells = grid.add_row().cells
        for cell, (label, value) in zip(cells, facts[start:start + cols]):
            label_par = cell.paragraphs[0]
            label_par.paragraph_format.space_after = Pt(0)
            label_run = label_par.add_run(label.upper())
            label_run.font.size = Pt(7.5)
            label_run.font.color.rgb = RGBColor(*_MUTED)
            value_par = cell.add_paragraph()
            value_par.paragraph_format.space_after = Pt(6)
            value_par.add_run(value).font.size = Pt(9.5)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Pt(6)


def render_docx(report: ReportModel) -> bytes:
    import docx
    from docx.shared import Inches, Pt, RGBColor

    doc = docx.Document()
    _style_document(doc)
    props = doc.core_properties
    props.title = report.title
    props.subject = (report.question or "")[:250]
    props.author = "Phronesis"
    props.comments = f"Run {report.run_id}" if report.run_id else "No stored run"
    props.keywords = "phronesis, provenance, report"

    _brand_header_footer(doc, report)
    _title_block(doc, report)
    for notice in report.notices:
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Pt(8)
        label = par.add_run("Notice. ")
        label.bold = True
        for run in (label, par.add_run(notice)):
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x92, 0x54, 0x00)

    if not question_is_title(report):
        doc.add_heading("Question", 1)
        doc.add_paragraph(report.question or "No question was supplied.")

    if report.task:
        doc.add_heading("Task status", 1)
        for line in task_lines(report.task):
            doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Answer", 1)
    add_markdown(doc, report.answer_markdown or "_No answer text._")

    sim = report.simulation
    if sim and (sim.rows or report.plot_png):
        doc.add_heading("Simulation", 1)
        label = " ".join(x for x in (sim.model_id, f"({sim.model_name})" if sim.model_name else "") if x)
        if label:
            line = f"Model: {label}"
            if sim.duration is not None:
                line += f", duration {fmt_number(sim.duration)} {sim.time_unit}".rstrip()
            if sim.convergence_status:
                line += f", {sim.convergence_status}"
            doc.add_paragraph(line)
        if report.plot_png:
            try:
                doc.add_picture(io.BytesIO(report.plot_png), width=Inches(6))
            except Exception:  # noqa: BLE001
                doc.add_paragraph("The simulation plot could not be embedded.")
        if sim.rows:
            unit = f" ({sim.quantity_unit})" if sim.quantity_unit else ""
            has_ss = any(r.steady_state is not None for r in sim.rows)
            header = ["Species", f"Initial{unit}", f"Final{unit}"] + ([f"Steady state{unit}"] if has_ss else [])
            rows = [
                [r.species, fmt_number(r.initial), fmt_number(r.final)] + ([fmt_number(r.steady_state)] if has_ss else [])
                for r in sim.rows
            ]
            _add_table(doc, header, rows, runs_mode=False)

    for table in report.data_tables:
        doc.add_heading(table.title, 1)
        _add_table(doc, table.header, table.rows, runs_mode=False)
        if table.note:
            doc.add_paragraph().add_run(table.note).italic = True

    doc.add_heading("Sources and citations", 1)
    if report.citations:
        for c in report.citations:
            doc.add_paragraph(f"[{c.n}] {citation_line(c)}")
    else:
        doc.add_paragraph("No structured evidence records accompany this answer.")

    doc.add_heading("Verification", 1)
    for line in verification_lines(report.verification):
        stripped = line.lstrip()
        style = "List Bullet 2" if line.startswith("  ") else "List Bullet"
        doc.add_paragraph(stripped[2:] if stripped.startswith("- ") else stripped, style=style)

    doc.add_heading("Evidence limitations", 1)
    if report.limitations:
        for lim in report.limitations:
            doc.add_paragraph(lim, style="List Bullet")
    else:
        doc.add_paragraph("No limitations were recorded.")

    doc.add_heading("Run metadata", 1)
    _add_table(doc, ["Field", "Value"], report.run_metadata, runs_mode=False)

    doc.add_heading("Reproducibility", 1)
    if report.tool_calls:
        _add_table(
            doc,
            ["#", "Tool", "Inputs", "Status", "Output sha256"],
            [[str(t.seq), t.tool, t.inputs, t.status, t.output_sha256] for t in report.tool_calls],
            runs_mode=False,
        )
    else:
        doc.add_paragraph("No tool calls were recorded for this answer.")
    for label, value in report.reproducibility:
        par = doc.add_paragraph(style="List Bullet")
        par.add_run(f"{label}: ").bold = True
        par.add_run(value)

    closing = doc.add_paragraph()
    closing.paragraph_format.space_before = Pt(12)
    closing_run = closing.add_run(report.footer)
    closing_run.italic = True
    closing_run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
