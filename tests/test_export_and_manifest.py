"""
Run manifest, report export and re-run diff (feature 9). Offline: tools are the recording fakes from
tests/fakes.py, upstream version/BioModels responses are trimmed live captures in
tests/fixtures/export/upstream_versions.json served through httpx.MockTransport.
"""

from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any, Dict

import docx  # python-docx, fpdf2 and pypdf are runtime dependencies: a missing one is a failure
import httpx
import pypdf
import pytest

from phronesis.agents.query_context import DatasetRef
from phronesis.provenance import run_manifest as rm
from phronesis.provenance import source_versions as sv
from phronesis.provenance.rerun import diff_runs, numeric_change
from phronesis.provenance.run_store import RunStore, get_run_store
from phronesis.tools.registry import AdapterType, ErrorType, ToolResult, ToolUniverseEntry
from phronesis.utils.memory.clinical_memory import ClinicalMemoryEngine
from tests import fakes
from tests.fakes import build_recording_registry, make_offline_orchestrator

FIXTURES = json.loads((Path(__file__).parent / "fixtures" / "export" / "upstream_versions.json").read_text(encoding="utf-8"))
SIM_QUERY = "Simulate BIOMD0000000010 over 120 seconds and summarize KRAS literature"
HEX64 = re.compile(r"^[0-9a-f]{64}$")


def _png_bytes() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 6), (40, 80, 160)).save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def runs_dir(monkeypatch, tmp_path):
    path = tmp_path / "runs"
    monkeypatch.setenv("PHRONESIS_RUNS_DIR", str(path))
    return path


@pytest.fixture
def offline(monkeypatch, tmp_path, runs_dir):
    return make_offline_orchestrator(monkeypatch, tmp_path)


@pytest.fixture
def with_plot(monkeypatch):
    """The fake Vivarix output gains a plot, as the real engine's does, to prove plots never reach the manifest."""
    original = fakes._fake_output
    png = base64.b64encode(_png_bytes()).decode()

    def fake(tool_name, kwargs):
        out = original(tool_name, kwargs)
        if tool_name == "vivarix_simulate_time_course":
            out = {**out, "plot_base64_png": png, "diagnostics": {"convergence_status": "DYNAMIC_EQUILIBRIUM"}}
        return out

    monkeypatch.setattr(fakes, "_fake_output", fake)


@pytest.fixture
def api(monkeypatch, tmp_path, runs_dir):
    from fastapi.testclient import TestClient

    from app.backend import main as backend_main

    engine = ClinicalMemoryEngine(storage_dir=str(tmp_path / "api_mem"))
    monkeypatch.setattr(backend_main, "get_clinical_memory_engine", lambda: engine)
    monkeypatch.delenv("PHRONESIS_AUTH_MODE", raising=False)
    monkeypatch.delenv("PHRONESIS_JWT_SECRET", raising=False)
    monkeypatch.setenv("PHRONESIS_API_TOKENS", "alice-token:alice,bob-token:bob")
    return TestClient(backend_main.app)


ALICE = {"Authorization": "Bearer alice-token"}
BOB = {"Authorization": "Bearer bob-token"}


def _calls_named(manifest: Dict[str, Any], name: str):
    return [c for c in manifest["tool_calls"] if c["tool_name"] == name]


# =============================================================================
# 1-5: the manifest
# =============================================================================

@pytest.mark.asyncio
async def test_manifest_records_validated_tool_calls(offline, with_plot, runs_dir):
    orch, _ = offline
    res = await orch.execute_inquiry(SIM_QUERY, thread_id="t-manifest", user_id="alice")
    m = res.run_manifest
    assert m and m["schema_version"] == "phronesis.run_manifest/1"

    viv = _calls_named(m, "vivarix_simulate_time_course")
    assert viv and viv[0]["inputs"]["biomodel_id"] == "BIOMD0000000010"
    # The interval count follows SimulationPlan (execution workstream); only its recording is checked here.
    assert viv[0]["inputs"]["duration"] == 120.0 and int(viv[0]["inputs"]["intervals"]) > 0
    assert viv[0]["output_digest"]["final_concentrations"] == {"Mos": 34.13, "Mos-P": 65.87}
    assert viv[0]["output_digest"]["convergence_status"] == "DYNAMIC_EQUILIBRIUM"
    lit = _calls_named(m, "lithoscribe_literature_synthesis")
    assert lit and lit[0]["inputs"]["max_articles"] == 8
    assert lit[0]["output_digest"]["pmids"] == ["123456"]

    seqs = [c["seq"] for c in m["tool_calls"]]
    assert seqs == sorted(seqs) and len(set(seqs)) == len(seqs) and seqs[0] == 1
    assert all(HEX64.match(c["output_sha256"]) for c in m["tool_calls"] if c["success"])
    text = json.dumps(m)
    assert "plot_base64_png" not in text
    assert len(text) < 15_000
    assert m["started_at"] <= m["finished_at"]
    assert m["response_sha256"] == hashlib.sha256(res.response.encode("utf-8")).hexdigest()
    assert m["user_id_hash"] == hashlib.sha256(b"alice").hexdigest()[:16] and "alice" not in m["user_id_hash"]
    assert m["environment"]["app_version"] and HEX64.match(m["environment"]["source_fingerprint"])
    assert m["synthesis"]["llm_ok"] is False and HEX64.match(m["synthesis"]["system_prompt_sha256"])
    assert set(m["verification"]) >= {"first_pass", "final", "correction_attempted"}

    stored = get_run_store().load("alice", m["run_id"])
    assert stored is not None and stored["run_id"] == m["run_id"]
    assert get_run_store().load_response("alice", m["run_id"]) == res.response
    assert (runs_dir / "alice" / f"{m['run_id']}.json").is_file()


@pytest.mark.asyncio
async def test_manifest_routing_and_steps(offline):
    orch, _ = offline
    res = await orch.execute_inquiry("Run a kinetic ODE simulation of the MAPK cascade for 60 seconds", thread_id="t-route")
    m = res.run_manifest
    steps = [s for s in m["steps"] if s["step"] == "biomodel_search"]
    assert steps and steps[0]["outputs"]["selected"] == "BIOMD0000000010"
    assert steps[0]["outputs"]["hits"] == ["BIOMD0000000010"] and steps[0]["outputs"]["fallback"] is False
    assert m["routing"]["is_simulation_intent"] is True and m["routing"]["sim_duration_s"] == 60.0
    assert set(m["memory_context"]) == {"n_evidence", "n_instructions", "memory_summary_sha256"}
    # Probes are off in tests: the pin says so instead of guessing a revision.
    pins = {p["accession"]: p for p in m["biomodels"]}
    assert pins["BIOMD0000000010"]["latest_revision"] is None and pins["BIOMD0000000010"]["error"]
    assert all(v["value"] is None and v["error"] for v in m["source_versions"].values())


@pytest.mark.asyncio
async def test_concurrent_runs_isolated(offline):
    orch, _ = offline
    a, b = await asyncio.gather(
        orch.execute_inquiry("What does the literature say about KRAS in pancreatic cancer?", thread_id="t-a"),
        orch.execute_inquiry("What does the literature say about EGFR in lung cancer?", thread_id="t-b"),
    )
    ma, mb = a.run_manifest, b.run_manifest
    assert ma["run_id"] != mb["run_id"]
    inputs_a = json.dumps([c["inputs"] for c in ma["tool_calls"]])
    inputs_b = json.dumps([c["inputs"] for c in mb["tool_calls"]])
    assert "KRAS" in inputs_a and "EGFR" not in inputs_a
    assert "EGFR" in inputs_b and "KRAS" not in inputs_b
    assert rm.current_recorder() is None


@pytest.mark.asyncio
async def test_registry_records_failures(runs_dir):
    reg, _ = build_recording_registry()

    async def slow(query: str, max_results: int = 3):
        await asyncio.sleep(1.0)
        return {"results": []}

    reg.register(
        ToolUniverseEntry(
            name="slow_tool", description="slow", domain="general", adapter_type=AdapterType.DIRECT_REST,
            input_schema=fakes._TOOL_SCHEMAS["websurfer_search"], timeout_seconds=0.05,
        ),
        executor=slow,
    )

    # Without an active recorder the registry is unaffected.
    assert rm.current_recorder() is None
    ok = await reg.execute("synapta_kegg_pathways", target_symbol="KRAS")
    assert ok.success

    rec = rm.RunRecorder(query="q", thread_id="t", model_name="m", enable_web_search=False, user_id="alice", api_key_supplied=False)
    rec.persist = False
    token = rm.activate_recorder(rec)
    try:
        bad = await reg.execute("vivarix_simulate_time_course", biomodel_id="BIOMD0000000010", duration=0.0)
        slow_res = await reg.execute("slow_tool", query="x")
        missing = await reg.execute("not_a_tool", foo=1)
        rm.record_tool_call(tool_name="weird", entry=None, inputs=None, raw_inputs={"x": object()}, result=None,
                            started_at=rm.utc_now(), elapsed_seconds=0.0)
    finally:
        rm.reset_recorder(token)
    assert bad.error_type == ErrorType.VALIDATION_ERROR and slow_res.error_type == ErrorType.TIMEOUT
    assert missing.error_type == ErrorType.NOT_FOUND

    m = await rec.finalize(response="", duration=0.1)
    by_name = {c["tool_name"]: c for c in m["tool_calls"]}
    assert by_name["vivarix_simulate_time_course"]["error_type"] == "VALIDATION_ERROR"
    assert by_name["vivarix_simulate_time_course"]["inputs"] == {}
    assert by_name["vivarix_simulate_time_course"]["raw_inputs"]["duration"] == 0.0
    assert by_name["slow_tool"]["error_type"] == "TIMEOUT" and by_name["slow_tool"]["timeout_seconds"] == 0.05
    assert by_name["not_a_tool"]["error_type"] == "NOT_FOUND"
    assert by_name["weird"]["success"] is False and by_name["weird"]["error"]
    assert "synapta_kegg_pathways" not in by_name
    assert not list(runs_dir.glob("**/*.json"))  # persist=False


@pytest.mark.asyncio
async def test_dataset_sha256(tmp_path):
    payload = b"AnnData-like bytes " * 1000
    h5ad = tmp_path / "nested" / "atlas.h5ad"
    h5ad.parent.mkdir()
    h5ad.write_bytes(payload)
    pin = rm.hash_dataset_path(str(h5ad))
    assert pin.sha256 == hashlib.sha256(payload).hexdigest()
    assert pin.path_basename == "atlas.h5ad" and pin.kind == "h5ad" and pin.size_bytes == len(payload)

    zarr = tmp_path / "cells.zarr"
    (zarr / "X").mkdir(parents=True)
    (zarr / ".zattrs").write_text("{}")
    (zarr / "X" / "0.0").write_bytes(b"\x00\x01\x02")
    first, second = rm.hash_dataset_path(str(zarr)), rm.hash_dataset_path(str(zarr))
    assert first.sha256 == second.sha256 and first.kind == "zarr" and first.note == "tree_hash"
    (zarr / "X" / "0.0").write_bytes(b"\x00\x01\x03")
    assert rm.hash_dataset_path(str(zarr)).sha256 != first.sha256

    assert rm.hash_dataset_path(str(h5ad), max_bytes=10).note == "not_hashed: too_large"
    assert rm.hash_dataset_path(str(tmp_path / "missing.h5ad")).note == "not_hashed: not_found"

    rec = rm.RunRecorder(query="q", thread_id="t", model_name="m", enable_web_search=False, user_id="u", api_key_supplied=False)
    rec.persist = False
    rec.start_dataset_hash(str(h5ad))
    ref = DatasetRef(id="ds_0123456789ab", kind="h5ad", display_name="upload.h5ad", analysis_path=str(h5ad),
                     public_provenance={"sha256": hashlib.sha256(payload).hexdigest(), "display_name": "upload.h5ad"})
    rec.start_dataset_hash(ref)
    m = await rec.finalize(response="", duration=0.0)
    pins = {d["path_basename"]: d for d in m["datasets"]}
    assert pins["atlas.h5ad"]["sha256"] == hashlib.sha256(payload).hexdigest()
    assert pins["upload.h5ad"]["sha256"] == hashlib.sha256(payload).hexdigest()
    assert str(tmp_path) not in json.dumps(m)


# =============================================================================
# 6: BioModels pin and data-source versions (recorded live responses)
# =============================================================================

def _transport(status_files=200, status_meta=200, html=False):
    files = FIXTURES["biomodels_files_BIOMD0000000010"]
    meta = FIXTURES["biomodels_meta_BIOMD0000000010"]
    not_found = FIXTURES["biomodels_files_BIOMD0000000010.3_404"]
    html_body = FIXTURES["biomodels_download_BIOMD0000000010.3_403_html_head"]

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "/model/files/" in url:
            status, body = status_files, files if status_files == 200 else not_found
        elif "/biomodels/" in url or "biomodels.org/BIOMD" in url:
            status, body = status_meta, meta if status_meta == 200 else not_found
        elif "opentargets" in url:
            status, body = 200, FIXTURES["opentargets_meta"]
        elif "kegg" in url:
            return httpx.Response(200, text=FIXTURES["kegg_info_pathway"], headers={"content-type": "text/plain; charset=utf-8"})
        elif "europepmc" in url:
            status, body = 200, FIXTURES["europepmc_search_10712587"]
        else:
            status, body = 404, {}
        if html and status != 200:
            return httpx.Response(status, text=html_body, headers={"content-type": "text/html;charset=ISO-8859-1"})
        return httpx.Response(status, json=body)

    return httpx.MockTransport(handler)


@pytest.mark.asyncio
async def test_biomodel_pin_resolver(monkeypatch, tmp_path):
    cache = tmp_path / "BIOMD0000000010.cps"
    cache.write_bytes(b"<COPASI cached model bytes>")
    monkeypatch.setattr(sv, "vivarix_cache_path", lambda acc: cache)

    pin = await sv.resolve_biomodel_pin("BIOMD0000000010", transport=_transport())
    assert pin["latest_revision"] == 4  # revisions are [1, 2, 4]: the last one, not the count
    assert pin["main_file"] == "BIOMD0000000010_url.xml"
    assert pin["main_file_sha256"] == "906964e07ffd570ad94a03c1222e230b54b913cd536b008392bfb6155930a966"
    assert pin["local_cache_sha256"] == hashlib.sha256(cache.read_bytes()).hexdigest()
    assert pin["error"] is None

    missing = await sv.resolve_biomodel_pin("BIOMD0000000010", transport=_transport(404, 404))
    assert missing["main_file_sha256"] is None and missing["latest_revision"] is None and "404" in missing["error"]
    forbidden = await sv.resolve_biomodel_pin("BIOMD0000000010", transport=_transport(403, 403, html=True))
    assert forbidden["main_file_sha256"] is None and "403" in forbidden["error"]
    assert (await sv.resolve_biomodel_pin("../etc/passwd", transport=_transport()))["error"] == "not a BioModels accession"

    versions = await sv.get_source_versions(transport=_transport())
    assert versions["opentargets"] == {"value": "26.09", "api": "26.9.0", "data": "26.09", "name": "Open Targets GraphQL & REST API Beta"}
    assert versions["kegg"]["value"] == "2026/10/09"
    assert versions["europepmc"]["value"] == "6.9"

    def down(request):
        raise httpx.ConnectError("offline", request=request)

    failed = await sv.get_source_versions(transport=httpx.MockTransport(down))
    assert all(v["value"] is None and "ConnectError" in v["error"] for v in failed.values())


# =============================================================================
# 7-10: sources and rendered reports
# =============================================================================

def _epmc_passage(**overrides):
    rec = FIXTURES["europepmc_search_10712587"]["resultList"]["result"][0]
    passage = {
        "passage_id": "p1", "pmid": rec["pmid"], "doi": rec["doi"], "title": rec["title"],
        "authors": rec["authorString"], "journal": rec["journalTitle"], "year": rec["pubYear"],
        "section": "Abstract", "passage_text": rec["title"], "has_full_text": False,
    }
    passage.update(overrides)
    return passage


def test_source_records_no_fabrication():
    from phronesis.extensions import hooks, load_extensions
    from phronesis.extensions.export import build_source_records

    load_extensions()
    assert any(getattr(fn, "__name__", "") == "_sources_artifact" for fn in hooks.get_artifact_builders())

    lit = ToolResult(success=True, tool_name="lithoscribe_literature_synthesis",
                     data={"ranked_passages": [_epmc_passage(year="", journal=""), _epmc_passage(passage_id="p2")]})
    web = ToolResult(success=True, tool_name="websurfer_search",
                     data={"results": [{"title": "EuropePMC record", "url": "https://europepmc.org/article/MED/10712587"}, {"title": "bad", "url": "javascript:alert(1)"}]})
    failed = ToolResult(success=False, tool_name="websurfer_search", error="down")
    records = build_source_records({"lithoscribe": lit, "websurfer": web})
    assert len(records) == 2  # the duplicate PMID collapses; the non-http URL is dropped
    first = records[0]
    assert first["pmid"] == "10712587" and first["year"] == "" and first["journal"] == ""
    assert first["url"] == "https://pubmed.ncbi.nlm.nih.gov/10712587/"
    assert records[1]["kind"] == "web" and records[1]["url"].startswith("https://europepmc.org/")
    assert "2024" not in json.dumps(records)
    assert build_source_records({"websurfer": failed}) == []


ANSWER_MD = """> ⚠️ **Scientific Verification Notice**: This synthesized assessment contains claims that could not be fully verified against empirical tool artifacts (Verification Score: 0.62 / 1.00).
> **Unverified Findings**: Mos is 50 nmol

### 🧬 Kinetic evidence for BIOMD0000000010

- Mos falls from 100.0 to 34.13 nmol over 120 s
  - Mos-P reaches 65.87 nmol, a rise ≥ 60 nmol at 5 µM scale
    1. nested ordered point
1. first ordered point

| Species | Initial (nmol) | Final (nmol) |
| --- | --- | --- |
| Mos | 100.0 | 34.13 |
| Mos-P | 0.0 | 65.87 |

Kholodenko 2000 (PMID: 10712587) describes the cascade [paper](https://doi.org/10.1046/j.1432-1327.2000.01197.x).

```text
final Mos = 34.13
```
"""


def _message(content: str = ANSWER_MD, manifest: Dict[str, Any] | None = None) -> Dict[str, Any]:
    return {
        "content": content,
        "model": "offline-test",
        "verification_report": {
            "passed": False, "verification_score": 0.62, "numerical_checks_passed": 2, "numerical_checks_failed": 1,
            "verified_identifiers": ["BIOMD0000000010", "PMID:10712587"],
            "defects": [{"category": "NUMERICAL", "severity": "CRITICAL", "description": "Mos value does not match the simulation", "claim": "Mos is 50 nmol"}],
        },
        "execution_artifacts": {
            "vivarix": {
                "model_id": "BIOMD0000000010", "model_name": "Fake MAPK", "duration": 120.0,
                "units": {"time": "s", "quantity": "nmol"},
                "initial_concentrations": {"Mos": 100.0, "Mos-P": 0.0},
                "final_concentrations": {"Mos": 34.13, "Mos-P": 65.87},
                "plot_base64_png": base64.b64encode(_png_bytes()).decode(),
            },
            "sources": [
                {"kind": "literature", "pmid": "10712587", "doi": "10.1046/j.1432-1327.2000.01197.x",
                 "title": FIXTURES["europepmc_search_10712587"]["resultList"]["result"][0]["title"],
                 "authors": "Kholodenko BN.", "journal": "", "year": "", "url": "https://pubmed.ncbi.nlm.nih.gov/10712587/",
                 "tool": "lithoscribe_literature_synthesis"},
            ],
        },
        "run_manifest": manifest,
    }


def test_render_docx_sections():
    from phronesis.export import build_report_model, render_docx

    report = build_report_model("What does BIOMD0000000010 predict for Mos?", _message())
    doc = docx.Document(io.BytesIO(render_docx(report)))
    texts = [p.text for p in doc.paragraphs]
    for section in ("Answer", "Simulation", "Sources and citations", "Verification", "Evidence limitations", "Run metadata", "Reproducibility"):
        assert section in texts
    # A short question is the title (no separate Question section); the header carries the brand mark and wordmark.
    assert next(p.text for p in doc.paragraphs if p.style.name == "Title") == "What does BIOMD0000000010 predict for Mos?"
    assert "Question" not in texts
    header = doc.sections[0].header.paragraphs[0]
    assert len(header._p.xpath(".//pic:pic")) == 2 and "Investigation report" in header.text
    assert "Page" in doc.sections[0].footer.paragraphs[0].text
    assert not any("Scientific Verification Notice" in t for t in texts)
    answer_table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Species" and "nmol" in t.rows[0].cells[1].text)
    assert [c.text for c in answer_table.rows[1].cells] == ["Mos", "100.0", "34.13"]
    assert len(doc.inline_shapes) == 1
    styles = {p.style.name for p in doc.paragraphs}
    assert {"List Bullet", "List Bullet 2", "List Number 3", "List Number"} <= styles
    assert any(r.font.name == "Consolas" for p in doc.paragraphs for r in p.runs)
    assert doc.styles["Normal"].font.name == "Google Sans" and doc.styles["Heading 1"].font.name == "Google Sans"
    assert any("(https://doi.org/10.1046/j.1432-1327.2000.01197.x)" in t for t in texts)
    assert doc.core_properties.title == "Phronesis Investigation Report"


def test_render_pdf_text():
    from phronesis.export import build_report_model, render_pdf

    report = build_report_model("🔬 What does BIOMD0000000010 predict? 📚", _message())
    data = render_pdf(report)
    assert data.startswith(b"%PDF")
    text = "\n".join(page.extract_text() for page in pypdf.PdfReader(io.BytesIO(data)).pages)
    assert "Kinetic evidence for BIOMD0000000010" in text
    assert "34.13" in text and "65.87" in text and "Mos-P" in text
    assert "≥" in text and "µ" in text
    assert "Scientific Verification Notice" not in text
    assert "Page 1 of" in text
    fonts = {
        str(f.get_object().get("/BaseFont", ""))
        for page in pypdf.PdfReader(io.BytesIO(data)).pages
        for f in (page["/Resources"].get("/Font") or {}).values()
    }
    assert any("GoogleSans" in name for name in fonts), fonts
    assert "What does BIOMD0000000010 predict?" in text  # the question is the title, without the stray emoji


def test_banner_stripped_and_section_present():
    from phronesis.export import build_report_model, render_markdown
    from phronesis.export.report_model import strip_verification_banner

    assert strip_verification_banner(ANSWER_MD).startswith("### 🧬 Kinetic evidence")
    assert strip_verification_banner("Plain answer") == "Plain answer"
    report = build_report_model("q", _message())
    assert "Scientific Verification Notice" not in report.answer_markdown
    md = render_markdown(report).decode("utf-8")
    assert md.startswith("**Phronesis** · Investigation report\n\n# Q\n")
    assert "simulation.png" not in md and "included in the PDF and Word exports" in md
    verification = md.split("## Verification", 1)[1].split("## Evidence limitations", 1)[0]
    assert "did not pass" in verification and "0.62 / 1.00" in verification
    assert "[CRITICAL] NUMERICAL: Mos value does not match the simulation (claim: Mos is 50 nmol)" in verification
    assert report.citations[0].cited_in_answer is True and report.citations[0].year == ""
    assert "Provenance: client supplied" in md


# =============================================================================
# 11-16: endpoints and re-runs
# =============================================================================

def _store_alice_run(orch, query=SIM_QUERY, **kwargs):
    return asyncio.run(orch.execute_inquiry(query, thread_id="t-alice", user_id="alice", **kwargs))


def test_export_endpoint_auth_and_types(api, offline):
    orch, _ = offline
    body = {"question": "What does BIOMD0000000010 predict for Mos?", "message": _message()}
    assert api.post("/api/export/docx", json=body).status_code == 401

    res = api.post("/api/export/docx", json=body, headers=ALICE)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    disposition = res.headers["content-disposition"]
    assert disposition.startswith("attachment;") and "phronesis-what-does-biomd0000000010-predict-for-mos-" in disposition
    assert res.content[:2] == b"PK"

    md = api.post("/api/export/md", json=body, headers=ALICE)
    assert md.status_code == 200 and md.headers["content-type"].startswith("text/markdown")
    as_json = api.post("/api/export/json", json=body, headers=ALICE)
    assert as_json.status_code == 200 and as_json.json()["report"]["provenance"] == "client_supplied"
    assert api.post("/api/export/exe", json=body, headers=ALICE).status_code == 422

    # A stored run is trusted; edited content is labelled.
    result = _store_alice_run(orch)
    stored_msg = _message(content=result.response, manifest=result.run_manifest)
    ok = api.post("/api/export/json", json={"question": SIM_QUERY, "message": stored_msg}, headers=ALICE).json()["report"]
    assert ok["provenance"] == "server_store" and ok["content_matches_run"] is True
    assert any(t["tool"] == "vivarix_simulate_time_course" for t in ok["tool_calls"])
    tampered = {**stored_msg, "content": result.response + "\nExtra claim."}
    bad = api.post("/api/export/json", json={"question": SIM_QUERY, "message": tampered}, headers=ALICE).json()["report"]
    assert bad["content_matches_run"] is False and "Content differs from the verified run output." in bad["notices"]
    as_bob = api.post("/api/export/json", json={"question": SIM_QUERY, "message": stored_msg}, headers=BOB).json()["report"]
    assert as_bob["provenance"] == "client_supplied"

    bundle = api.post("/api/export/bundle", json={"question": SIM_QUERY, "message": stored_msg}, headers=ALICE)
    assert bundle.status_code == 200 and bundle.headers["content-type"] == "application/zip"
    names = set(zipfile.ZipFile(io.BytesIO(bundle.content)).namelist())
    assert {"report.pdf", "report.docx", "answer.md", "manifest.json", "sources.json", "REPRODUCE.md"} <= names


def test_rerun_tools_mode_identical(api, offline):
    orch, _ = offline
    run_id = _store_alice_run(orch).run_manifest["run_id"]
    res = api.post(f"/api/runs/{run_id}/rerun", json={"mode": "tools"}, headers=ALICE)
    assert res.status_code == 200
    report = res.json()
    assert report["overall"] == "identical" and report["mode"] == "tools" and report["provenance"] == "server_store"
    assert report["tool_diffs"] and all(d["output_hash_equal"] and not d["changes"] for d in report["tool_diffs"])
    assert report["missing_in_new"] == [] and report["new_only"] == []
    assert report["new_run_id"] != run_id
    # Replays are diffed, not stored as runs of their own.
    assert get_run_store().load("alice", report["new_run_id"]) is None


def test_rerun_detects_changes(api, offline, monkeypatch):
    orch, _ = offline
    run_id = _store_alice_run(orch).run_manifest["run_id"]
    original = fakes._fake_output

    def drifted(tool_name, kwargs):
        out = original(tool_name, kwargs)
        if tool_name == "vivarix_simulate_time_course":
            out = {**out, "final_concentrations": {"Mos": 40.0, "Mos-P": 65.87}}
        if tool_name == "lithoscribe_literature_synthesis":
            out = {"ranked_passages": [_epmc_passage()]}
        return out

    monkeypatch.setattr(fakes, "_fake_output", drifted)
    report = api.post(f"/api/runs/{run_id}/rerun", json={"mode": "tools"}, headers=ALICE).json()
    assert report["overall"] == "changed"
    changes = [c for d in report["tool_diffs"] for c in d["changes"]]
    mos = next(c for c in changes if c["path"] == "final_concentrations.Mos")
    assert mos["tier"] == "changed" and mos["before"] == 34.13 and mos["after"] == 40.0
    assert mos["rel_change"] == pytest.approx(0.172, abs=1e-3)
    assert any(c["path"] == "pmids" and c["before"] == "123456" and c["tier"] == "removed" for c in changes)
    assert any(c["path"] == "pmids" and c["after"] == "10712587" and c["tier"] == "added" for c in changes)
    assert not any(c["path"] == "final_concentrations.Mos-P" for c in changes)


def test_numeric_tiers():
    assert numeric_change("x", 280.0, 280.0).tier == "identical"
    assert numeric_change("x", 280.0, 281.3).tier == "within_tolerance"
    assert numeric_change("x", 34.13, 40.0).tier == "changed"
    assert numeric_change("x", 0.0, 1.0).tier == "changed"
    base = {"run_id": "a", "tool_calls": [], "source_versions": {"kegg": {"value": "2026/10/09"}}}
    newer = {"run_id": "b", "tool_calls": [], "source_versions": {"kegg": {"value": "2026/11/01"}}}
    assert diff_runs(base, newer).overall == "drifted"


def test_rerun_scoped_to_user(api, offline):
    orch, _ = offline
    run_id = _store_alice_run(orch).run_manifest["run_id"]
    assert api.get(f"/api/runs/{run_id}", headers=ALICE).status_code == 200
    assert api.get(f"/api/runs/{run_id}", headers=BOB).status_code == 404
    assert api.post(f"/api/runs/{run_id}/rerun", json={"mode": "tools"}, headers=BOB).status_code == 404
    assert api.get("/api/runs/..%2F..%2Fsecrets", headers=ALICE).status_code == 404
    assert api.get(f"/api/runs/{run_id}").status_code == 401
    assert "_response_text" not in api.get(f"/api/runs/{run_id}", headers=ALICE).json()


def test_client_manifest_rejects_unknown_tool(api, offline):
    orch, _ = offline
    run_id = _store_alice_run(orch).run_manifest["run_id"]
    manifest = api.get(f"/api/runs/{run_id}", headers=ALICE).json()

    ok = api.post("/api/runs/rerun", json={"mode": "tools", "manifest": manifest}, headers=BOB)
    assert ok.status_code == 200 and ok.json()["provenance"] == "client_supplied" and ok.json()["overall"] == "identical"

    evil = json.loads(json.dumps(manifest))
    evil["tool_calls"][0]["tool_name"] = "os_system"
    assert api.post("/api/runs/rerun", json={"mode": "tools", "manifest": evil}, headers=ALICE).status_code == 400
    upload = json.loads(json.dumps(manifest))
    upload["tool_calls"][0]["inputs"]["sbml_path"] = "C:/server/uploads/model.xml"
    assert api.post("/api/runs/rerun", json={"mode": "tools", "manifest": upload}, headers=ALICE).status_code == 400
    assert api.post("/api/runs/rerun", json={"mode": "tools", "manifest": {"run_id": "x"}}, headers=ALICE).status_code == 422


def test_full_rerun_does_not_persist_memory(api, offline, monkeypatch):
    from app.backend.routers import export as export_router

    orch, _ = offline
    monkeypatch.setattr(export_router, "_get_orchestrator", lambda: orch)
    run_id = _store_alice_run(orch).run_manifest["run_id"]
    before = orch.memory.get_memory_stats(user_id="alice")

    res = api.post(f"/api/runs/{run_id}/rerun", json={"mode": "full"}, headers=ALICE)
    assert res.status_code == 200
    report = res.json()
    assert report["mode"] == "full" and report["verification_delta"] is not None
    assert report["answer_unified_diff"] is not None
    assert orch.memory.get_memory_stats(user_id="alice") == before


def test_run_store_rejects_traversal(tmp_path):
    store = RunStore(tmp_path)
    with pytest.raises(ValueError):
        store.save("alice", {"run_id": "../../evil"})
    with pytest.raises(ValueError):
        store.save("../alice", {"run_id": "0" * 32})
    path = store.save("org:alice", {"run_id": "a" * 32, "x": 1})
    assert path.parent.parent == tmp_path and ":" not in path.parent.name
    assert store.load("org:alice", "a" * 32) == {"run_id": "a" * 32, "x": 1}
    assert store.load("bob", "a" * 32) is None


# =============================================================================
# Review regressions
# =============================================================================

def test_pdf_never_loads_markdown_images(tmp_path):
    from phronesis.export import build_report_model, render_pdf
    from phronesis.export.markdown_blocks import render_html

    local = tmp_path / "server_secret.png"
    local.write_bytes(_png_bytes())
    content = f"Plot ![secret plot]({local.as_posix()}) and ![remote](https://example.invalid/x.png)"
    assert "<img" not in render_html(content)
    data = render_pdf(build_report_model("q", {"content": content}))
    reader = pypdf.PdfReader(io.BytesIO(data))
    # Only the report's own brand images, never the ones the answer links to.
    branding = pypdf.PdfReader(io.BytesIO(render_pdf(build_report_model("q", {"content": "Plot"}))))
    assert sum(len(p.images) for p in reader.pages) == sum(len(p.images) for p in branding.pages)
    assert "[secret plot]" in "\n".join(p.extract_text() for p in reader.pages)


def test_renderers_survive_control_chars_and_corrupt_plot():
    from phronesis.export import build_report_model, render_docx, render_pdf

    corrupt = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"truncated").decode()
    msg = {
        "content": "Mos \x0bfalls\x00 to 34.13",
        "execution_artifacts": {"vivarix": {"plot_base64_png": corrupt, "final_concentrations": {"Mos": 34.13}}},
    }
    report = build_report_model("q\x1b", msg)
    assert report.plot_png is None and report.answer_markdown == "Mos falls to 34.13" and report.question == "q"
    assert render_docx(report)[:2] == b"PK" and render_pdf(report).startswith(b"%PDF")


def test_sources_dedupe_and_carry_retraction():
    from phronesis.export import build_report_model
    from phronesis.extensions.export import build_source_records
    from phronesis.sources import CitationRecord, EvidenceSource, register_source, unregister_source

    class _Grading(EvidenceSource):
        key = "review_fake_grading"

        def citations(self, result):
            return [CitationRecord(type="pmid", id="10712587", retracted=True)]

    register_source(_Grading())
    try:
        lit = ToolResult(success=True, tool_name="lithoscribe_literature_synthesis",
                         data={"ranked_passages": [_epmc_passage(retracted=None)]})
        graded = ToolResult(success=True, tool_name="fake", data={})
        records = build_source_records({"lithoscribe": lit, "review_fake_grading": graded})
    finally:
        unregister_source("review_fake_grading")
    assert [r["pmid"] for r in records] == ["10712587"] and records[0]["retracted"] is True

    msg = _message()
    msg["citations"] = {"entries": {"PMID:10712587": {"type": "pmid", "id": "10712587", "verified": True, "url": "https://europepmc.org/abstract/MED/10712587"}}}
    msg["execution_artifacts"]["sources"] = records
    report = build_report_model("q", msg)
    assert len(report.citations) == 1 and report.citations[0].retracted is True
    assert "RETRACTED" in __import__("phronesis.export.report_model", fromlist=["x"]).citation_line(report.citations[0])


def test_export_flags_posted_values_that_differ_from_stored_run(api, offline):
    orch, _ = offline
    result = _store_alice_run(orch)
    msg = _message(content=result.response, manifest=result.run_manifest)
    msg["execution_artifacts"] = result.execution_artifacts
    msg["verification_report"] = result.verification_report
    ok = api.post("/api/export/json", json={"question": SIM_QUERY, "message": msg}, headers=ALICE).json()["report"]
    assert ok["provenance"] == "server_store" and ok["notices"] == []

    edited = json.loads(json.dumps(msg))
    edited["execution_artifacts"]["vivarix"]["final_concentrations"]["Mos"] = 99.0
    edited["verification_report"]["verification_score"] = 0.99
    edited["verification_report"]["passed"] = not result.verification_report["passed"]
    bad = api.post("/api/export/json", json={"question": SIM_QUERY, "message": edited}, headers=ALICE).json()["report"]
    assert any("simulation values differ" in n for n in bad["notices"])
    assert any("verification report differs" in n for n in bad["notices"])

    junk = {**msg, "run_manifest": {"tool_calls": ["not-a-call"], "synthesis": "x"}}
    assert api.post("/api/export/md", json={"question": "q", "message": junk}, headers=ALICE).status_code in (200, 422)


def test_client_view_redacts_paths_in_errors():
    manifest = {"tool_calls": [{
        "tool_name": "vivarix_simulate_time_course", "success": False,
        "inputs": {"biomodel_id": "ds_0123456789ab", "sbml_path": r"C:\srv\uploads\alice\model.xml"},
        "error": r"Could not load C:\srv\uploads\alice\model.xml (C:/srv/uploads/alice/model.xml)",
    }, "garbage"]}
    view = rm.client_view(manifest)
    call = view["tool_calls"][0]
    assert call["inputs"]["sbml_path"] == "model.xml"
    assert "uploads" not in call["error"] and call["error"].count("model.xml") == 2


def test_run_store_separates_case_and_trailing_dot_ids(tmp_path):
    store = RunStore(tmp_path)
    run = {"run_id": "b" * 32}
    paths = {store.save(uid, run).parent.name.lower() for uid in ("alice", "Alice", "alice.")}
    assert len(paths) == 3
    assert store.load("Alice", "b" * 32) is not None and store.load("ALICE", "b" * 32) is None
